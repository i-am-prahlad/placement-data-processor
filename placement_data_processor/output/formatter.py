"""Consistent Excel formatting. Does not modify source workbooks."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.worksheet.errors import IgnoredError, IgnoredErrors
from openpyxl.worksheet.worksheet import Worksheet

# Ensure sqref is serialised in IgnoredError
if "sqref" not in IgnoredError.__attrs__:
    IgnoredError.__attrs__ = ("sqref",) + IgnoredError.__attrs__

# Patch WorksheetWriter.write_tail to serialize ignoredErrors element in worksheet XML
_orig_write_tail = WorksheetWriter.write_tail


def _patched_write_tail(self):
    self.write_protection()
    self.write_scenarios()
    self.write_filter()
    self.write_merged_cells()
    self.write_formatting()
    self.write_validations()
    self.write_hyperlinks()
    self.write_print()
    self.write_margins()
    self.write_page()
    self.write_header()
    self.write_breaks()
    if getattr(self.ws, "ignored_errors", None):
        self.xf.send(self.ws.ignored_errors.to_tree())
    self.write_drawings()
    self.write_legacy()
    self.write_tables()


WorksheetWriter.write_tail = _patched_write_tail

FONT_NAME = "Calibri"
FONT_SIZE = 11
MAX_COL_WIDTH = 42
MIN_COL_WIDTH = 10
ROW_HEIGHT = 14.4


BLACK_THIN = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)

THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)

# Header for RAIT & Remove: No fill, Bold, Black text
RAIT_HEADER_FONT = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color="000000")
RAIT_BODY_FONT = Font(name=FONT_NAME, size=FONT_SIZE, bold=False, color="000000")
# Light red fill only for College Roll No. data cells
LIGHT_RED_FILL = PatternFill("solid", fgColor="FFC7CE")

# General center alignment without wrapping
NO_WRAP_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=False)

# Legacy styles for Audit and Run Info
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT_NAME, size=FONT_SIZE, bold=False, color="000000")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WHITE = PatternFill("solid", fgColor="FFFFFF")


def is_resume_url(value: str) -> bool:
    text = (value or "").strip().lower()
    return text.startswith("http://") or text.startswith("https://")


def auto_fit_columns(ws: Worksheet) -> None:
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        longest = 0
        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row, col)
            val = cell.value
            if val is not None:
                # If merged cell or section title spans all columns, skip inflating column width
                if len(str(val)) > 80:
                    continue
                longest = max(longest, len(str(val)))
        ws.column_dimensions[letter].width = max(MIN_COL_WIDTH, longest + 3)


def style_rait_sheet(ws: Worksheet, roll_no_col_idx: int = 2, mobile_col_idx: int = 5) -> None:
    """Style RAIT sheet per Changes.md requirements."""
    if ws.max_row < 1 or ws.max_column < 1:
        return

    # Apply text format to entire mobile column and ignore NumberStoredAsText error
    mobile_col_letter = get_column_letter(mobile_col_idx)
    ws.column_dimensions[mobile_col_letter].number_format = "@"
    ws.ignored_errors = IgnoredErrors(ignoredError=[
        IgnoredError(sqref=f"{mobile_col_letter}1:{mobile_col_letter}1048576", numberStoredAsText=True)
    ])

    # Header row (Row 1)
    ws.row_dimensions[1].height = ROW_HEIGHT
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(1, col)
        cell.font = RAIT_HEADER_FONT
        cell.alignment = NO_WRAP_CENTER
        cell.border = BLACK_THIN
        # No fill on header

    # Populated data rows
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = ROW_HEIGHT
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            cell.font = RAIT_BODY_FONT
            cell.alignment = NO_WRAP_CENTER
            cell.border = BLACK_THIN

            # Store mobile numbers as clean text
            if col == mobile_col_idx:
                cell.number_format = "@"
                cell.data_type = "s"

            # Light red fill ONLY on College Roll No. data cells
            if col == roll_no_col_idx and cell.value:
                cell.fill = LIGHT_RED_FILL

            if cell.hyperlink:
                cell.font = Font(name=FONT_NAME, size=FONT_SIZE, color="0563C1", underline="single")

    auto_fit_columns(ws)
    # Ensure default gridlines are shown and NO filters exist
    ws.sheet_view.showGridLines = True
    ws.auto_filter.ref = None
    ws.freeze_panes = "A2"


def style_remove_sheet(ws: Worksheet, roll_no_col_idx: int = 2, mobile_col_idx: int = 5) -> None:
    """Style Remove sheet per Changes.md requirements."""
    if ws.max_row < 1 or ws.max_column < 1:
        return

    # Apply text format to entire mobile column and ignore NumberStoredAsText error
    mobile_col_letter = get_column_letter(mobile_col_idx)
    ws.column_dimensions[mobile_col_letter].number_format = "@"
    ws.ignored_errors = IgnoredErrors(ignoredError=[
        IgnoredError(sqref=f"{mobile_col_letter}1:{mobile_col_letter}1048576", numberStoredAsText=True)
    ])

    for row in range(1, ws.max_row + 1):
        ws.row_dimensions[row].height = ROW_HEIGHT
        # Check if row is a section title or company callout
        first_val = ws.cell(row, 1).value
        is_section_header = isinstance(first_val, str) and (
            first_val.startswith("---") or first_val.startswith("Company:") or "Section" in first_val
        )
        is_table_header = (first_val == "Sr No")

        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            if cell.value is None and is_section_header and col > 1:
                continue

            if is_table_header:
                cell.font = RAIT_HEADER_FONT
                cell.alignment = NO_WRAP_CENTER
                cell.border = BLACK_THIN
            elif is_section_header:
                cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color="000000")
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
            else:
                cell.font = RAIT_BODY_FONT
                cell.alignment = NO_WRAP_CENTER
                cell.border = BLACK_THIN
                if col == mobile_col_idx:
                    cell.number_format = "@"
                    cell.data_type = "s"
                if col == roll_no_col_idx and cell.value:
                    cell.fill = LIGHT_RED_FILL
                if cell.hyperlink:
                    cell.font = Font(name=FONT_NAME, size=FONT_SIZE, color="0563C1", underline="single")

    auto_fit_columns(ws)
    ws.sheet_view.showGridLines = True
    ws.auto_filter.ref = None



def style_worksheet(ws: Worksheet) -> None:
    """Legacy styling for Audit and Run Info worksheets."""
    if ws.max_row < 1 or ws.max_column < 1:
        return

    for col in range(1, ws.max_column + 1):
        cell = ws.cell(1, col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN

    for row in range(2, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            cell.font = BODY_FONT
            cell.alignment = CENTER
            cell.border = THIN
            cell.fill = WHITE
            if cell.hyperlink:
                cell.font = Font(name=FONT_NAME, size=FONT_SIZE, color="0563C1", underline="single")

    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        longest = 0
        for row in range(1, ws.max_row + 1):
            value = ws.cell(row, col).value
            longest = max(longest, len(str(value)) if value is not None else 0)
        ws.column_dimensions[letter].width = min(MAX_COL_WIDTH, max(MIN_COL_WIDTH, longest + 2))

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 22
