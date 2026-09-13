"""Build the final Excel workbook: RAIT, Remove, Audit, Run Info."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from config.output_schema import AUDIT_COLUMNS, RAIT_COLUMNS, REMOVE_COLUMNS, RUN_INFO_FIELDS
from config.settings import DriveRules
from core.processor import ProcessResult
from output.formatter import is_resume_url, style_rait_sheet, style_remove_sheet, style_worksheet
from utils.phones import format_mobile_text


RESUME_COL_INDEX = RAIT_COLUMNS.index("Resume") + 1
ROLL_NO_COL_INDEX = RAIT_COLUMNS.index("College Roll No.") + 1
MOBILE_COL_INDEX = RAIT_COLUMNS.index("Mobile Number") + 1


def _numeric_or_empty(value):
    """Convert numeric values to float or int for proper Excel numeric storage, or empty string if None/NaN."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value
    if isinstance(value, str):
        val_str = value.strip().replace("%", "")
        try:
            num = float(val_str)
            if num.is_integer():
                return int(num)
            return num
        except ValueError:
            return value
    return value


def _text_or_empty(value) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _student_dict_to_row(rec: dict | pd.Series, sr_no: int, include_job_role: bool) -> dict:
    norm_m = format_mobile_text(rec.get("mobile")) if hasattr(rec, "get") else ""
    raw_m = format_mobile_text(rec.get("raw_mobile")) if hasattr(rec, "get") else ""
    mobile = norm_m if norm_m else raw_m
    job_role = rec.get("job_role", "") if include_job_role else ""

    return {
        "Sr No": sr_no,
        "College Roll No.": _text_or_empty(rec.get("roll_no")),
        "Name": _text_or_empty(rec.get("name")),
        "Email Address": _text_or_empty(rec.get("email")),
        "Mobile Number": mobile or "",
        "Gender": _text_or_empty(rec.get("gender")),
        "Branch": _text_or_empty(rec.get("branch")),
        "College Name": _text_or_empty(rec.get("college")),
        "Job Role": job_role,
        "10th Percentage": _numeric_or_empty(rec.get("tenth_percentage")),
        "10th Year of Passing": _numeric_or_empty(rec.get("tenth_year")),
        "12th/Diploma?": _text_or_empty(rec.get("twelfth_or_diploma")),
        "12th/Diploma Percentage": _numeric_or_empty(rec.get("twelfth_percentage")),
        "12th Year of Passing": _numeric_or_empty(rec.get("twelfth_year")),
        "BE CGPA": _numeric_or_empty(rec.get("be_cgpa")),
        "BE Percentage": _numeric_or_empty(rec.get("be_percentage")),
        "MTech CGPA": _numeric_or_empty(rec.get("mtech_cgpa")),
        "MTech Percentage": _numeric_or_empty(rec.get("mtech_percentage")),
        "MTech Year of Passing": _numeric_or_empty(rec.get("mtech_year")),
        "Resume": _text_or_empty(rec.get("resume")),
    }


def build_rait_frame(eligible: pd.DataFrame, rules: DriveRules | None) -> pd.DataFrame:
    include_job_role = True if rules is None else rules.output.include_job_role
    ordered = eligible.sort_values(by=["name", "roll_no"], kind="mergesort").reset_index(drop=True)
    rows = []
    for i, rec in ordered.iterrows():
        rows.append(_student_dict_to_row(rec, i + 1, include_job_role))
    if not rows:
        return pd.DataFrame(columns=RAIT_COLUMNS)
    return pd.DataFrame(rows)[RAIT_COLUMNS]


def _write_sheet(ws, frame: pd.DataFrame, columns: list[str], hyperlink_col: int | None = None) -> None:
    mobile_col_idx = (columns.index("Mobile Number") + 1) if "Mobile Number" in columns else None
    if mobile_col_idx:
        col_letter = get_column_letter(mobile_col_idx)
        ws.column_dimensions[col_letter].number_format = "@"

    ws.append(columns)
    for _, row in frame.iterrows():
        values = [row.get(col, "") for col in columns]
        ws.append(values)
        curr_row = ws.max_row
        if mobile_col_idx:
            cell = ws.cell(curr_row, mobile_col_idx)
            val = "" if cell.value is None else str(cell.value)
            cell.value = val
            cell.data_type = "s"
            cell.number_format = "@"
        if hyperlink_col:
            cell = ws.cell(curr_row, hyperlink_col)
            url = str(cell.value or "").strip()
            if is_resume_url(url):
                cell.hyperlink = url
                cell.value = url
                cell.style = "Hyperlink"



def _write_remove_sections(ws, result: ProcessResult) -> None:
    """
    Write the Remove sheet divided into 6 distinct sections in this exact order:
    1. Placed Student
    2. Internship Students
    3. Removed due to 10th Percentage
    4. Removed due to 12th Percentage
    5. Removed due to BE CGPA
    6. Removed due to KT

    For Placed and Internship students: first mention company/companies, then student details.
    All student rows match the exact RAIT schema.
    """
    include_job_role = True if result.rules is None else result.rules.output.include_job_role
    logger_rows = result.removal_logger.rows if result.removal_logger else []

    # Map logger rows to the 6 sections
    def _categorize(item: dict) -> str | None:
        cat = item.get("Removal Category", "")
        if cat in {"Single placed", "Double placed", "Placement"}:
            return "Placed Student"
        if cat == "Internship":
            return "Internship Students"
        if cat == "10th eligibility":
            return "Removed due to 10th Percentage"
        if cat == "12th eligibility":
            return "Removed due to 12th Percentage"
        if cat == "CGPA eligibility":
            return "Removed due to BE CGPA"
        if cat == "Live KT / backlog":
            return "Removed due to KT"
        return None

    section_buckets: dict[str, list[dict]] = {
        "Placed Student": [],
        "Internship Students": [],
        "Removed due to 10th Percentage": [],
        "Removed due to 12th Percentage": [],
        "Removed due to BE CGPA": [],
        "Removed due to KT": [],
    }

    for item in logger_rows:
        sec = _categorize(item)
        if sec:
            section_buckets[sec].append(item)

    sections_order = [
        "Placed Student",
        "Internship Students",
        "Removed due to 10th Percentage",
        "Removed due to 12th Percentage",
        "Removed due to BE CGPA",
        "Removed due to KT",
    ]

    mobile_col_idx = REMOVE_COLUMNS.index("Mobile Number") + 1
    col_letter = get_column_letter(mobile_col_idx)
    ws.column_dimensions[col_letter].number_format = "@"

    for sec_idx, sec_name in enumerate(sections_order):
        items = section_buckets[sec_name]

        # Section Header banner
        ws.append([f"--- Section: {sec_name} ---"])
        ws.append(REMOVE_COLUMNS)

        if not items:
            ws.append([])  # Blank spacing line
            continue

        sr_counter = 1
        for item in items:
            rec = item.get("_student_record") or {}
            student_row = _student_dict_to_row(rec, sr_counter, include_job_role)
            sr_counter += 1

            # Placement or Internship company callout before student details
            if sec_name in {"Placed Student", "Internship Students"}:
                comp = item.get("Company", "") or rec.get("placement_companies", "") or rec.get("internship_company", "") or "Not Specified"
                label = "Company Placed at:" if sec_name == "Placed Student" else "Internship Company:"
                ws.append([f"{label} {comp}"])

            values = [student_row.get(col, "") for col in REMOVE_COLUMNS]
            ws.append(values)
            curr_row = ws.max_row
            if mobile_col_idx:
                cell = ws.cell(curr_row, mobile_col_idx)
                val = "" if cell.value is None else str(cell.value)
                cell.value = val
                cell.data_type = "s"
                cell.number_format = "@"
            if RESUME_COL_INDEX:
                cell = ws.cell(curr_row, RESUME_COL_INDEX)
                url = str(cell.value or "").strip()
                if is_resume_url(url):
                    cell.hyperlink = url
                    cell.value = url
                    cell.style = "Hyperlink"

        ws.append([])  # Blank spacing line between sections


def _write_run_info(ws, summary_map: dict) -> None:
    ws.append(["Field", "Value"])
    for field in RUN_INFO_FIELDS:
        ws.append([field, summary_map.get(field, "")])


def build_workbook(result: ProcessResult) -> bytes:
    rait = build_rait_frame(result.eligible, result.rules)
    audit = result.audit
    if audit is None or audit.empty:
        audit = pd.DataFrame(columns=AUDIT_COLUMNS)
    else:
        audit = audit[AUDIT_COLUMNS]

    wb = Workbook()
    ws_rait = wb.active
    ws_rait.title = "RAIT"
    _write_sheet(ws_rait, rait, RAIT_COLUMNS, hyperlink_col=RESUME_COL_INDEX)

    ws_remove = wb.create_sheet("Remove")
    _write_remove_sections(ws_remove, result)

    ws_audit = wb.create_sheet("Audit")
    _write_sheet(ws_audit, audit, AUDIT_COLUMNS)

    ws_info = wb.create_sheet("Run Info")
    _write_run_info(ws_info, result.summary.as_run_info())

    style_rait_sheet(ws_rait, roll_no_col_idx=ROLL_NO_COL_INDEX, mobile_col_idx=MOBILE_COL_INDEX)
    style_remove_sheet(ws_remove, roll_no_col_idx=ROLL_NO_COL_INDEX, mobile_col_idx=MOBILE_COL_INDEX)
    style_worksheet(ws_audit)

    style_worksheet(ws_info)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

