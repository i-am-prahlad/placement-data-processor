from io import BytesIO

from openpyxl import load_workbook

from config.output_schema import RAIT_COLUMNS, REMOVE_SECTIONS

from core.processor import process
from output.workbook import build_rait_frame, build_workbook
from tests.conftest import excel_bytes, open_rules, placement_row, response
from utils.names import to_proper_case
from utils.phones import normalize_mobile


def test_rait_column_order():
    responses = excel_bytes([response(), response(**{"College Roll No.": "R002", "Full Name": "ZED ZULU", "Email address": "z@x.com", "Mobile No": "9123456789"})])
    placement = excel_bytes([placement_row("R001"), placement_row("R002")])
    _, result = process(responses, placement, open_rules())
    frame = build_rait_frame(result.eligible, result.rules)
    assert list(frame.columns) == RAIT_COLUMNS
    assert frame.columns[1] == "College Roll No."


def test_sorting_by_name():
    responses = excel_bytes(
        [
            response(**{"College Roll No.": "R002", "Full Name": "ZARA KHAN", "Email address": "z@x.com", "Mobile No": "9123456789"}),
            response(**{"College Roll No.": "R001", "Full Name": "AMIT SHAH", "Email address": "a@x.com", "Mobile No": "9123456780"}),
        ]
    )
    placement = excel_bytes([placement_row("R001"), placement_row("R002")])
    _, result = process(responses, placement, open_rules())
    frame = build_rait_frame(result.eligible, result.rules)
    assert list(frame["Name"]) == ["Amit Shah", "Zara Khan"]
    assert list(frame["Sr No"]) == [1, 2]


def test_proper_case_and_mobile_and_hyperlink():
    assert to_proper_case("PRAHLAD SINGH PANNU") == "Prahlad Singh Pannu"
    assert normalize_mobile("+91 98765 43210") == ("9876543210", None)
    assert normalize_mobile("09876543210") == ("9876543210", None)
    assert normalize_mobile("123")[0] is None

    responses = excel_bytes(
        [response(**{"Full Name": "PRAHLAD SINGH PANNU", "Resume": "https://drive.google.com/file/d/abc", "Mobile No": "+91-9876543210"})]
    )
    placement = excel_bytes([placement_row("R001")])
    _, result = process(responses, placement, open_rules())
    payload = build_workbook(result)
    wb = load_workbook(BytesIO(payload))
    assert wb.sheetnames == ["RAIT", "Remove", "Audit", "Run Info"]
    rait = wb["RAIT"]
    assert [cell.value for cell in rait[1]] == RAIT_COLUMNS
    assert rait["B2"].value == "R001"
    # Light red fill on Roll No data cell
    assert rait["B2"].fill.fill_type == "solid"
    assert "FFC7CE" in str(rait["B2"].fill.fgColor.value).upper() or "FFCCCC" in str(rait["B2"].fill.fgColor.value).upper()
    # Header row has no fill, bold, black text
    assert rait["A1"].font.name == "Calibri"
    assert rait["A1"].font.bold is True
    assert rait["A1"].font.color.value in {"000000", "00000000", None}
    assert rait["A1"].fill.fill_type is None
    # No filters on RAIT header
    assert rait.auto_filter.ref is None
    # Black borders on populated cells
    assert rait["A2"].border.left.color.value in {"000000", "00000000"}

    # Name is in column C
    assert rait["C2"].value == "Prahlad Singh Pannu"
    assert rait["E2"].value == "9876543210"
    # Resume is in column T (col 20)
    assert rait.cell(2, 20).hyperlink is not None
    assert "http" in (rait.cell(2, 20).hyperlink.target or "")
    # Numeric values are numeric in Excel
    assert isinstance(rait.cell(2, 10).value, (int, float))  # 10th percentage
    assert isinstance(rait.cell(2, 15).value, (int, float))  # BE CGPA

    # Row height is <= 15
    assert rait.row_dimensions[1].height <= 15
    assert rait.row_dimensions[2].height <= 15


def test_remove_sections_and_formatting():
    responses = excel_bytes([
        response(**{"College Roll No.": "R001", "Full Name": "ALICE", "10th Percentage": 50}),
        response(**{"College Roll No.": "R002", "Full Name": "BOB", "12th/Diploma Percentage": 50}),
    ])
    placement = excel_bytes([placement_row("R001"), placement_row("R002")])
    rules = open_rules(academic={"mode": "threshold", "tenth": 70, "twelfth": 70, "cgpa": 7.0})
    _, result = process(responses, placement, rules)
    payload = build_workbook(result)
    wb = load_workbook(BytesIO(payload))
    ws_remove = wb["Remove"]
    assert ws_remove.auto_filter.ref is None

    text_contents = []
    for r in range(1, ws_remove.max_row + 1):
        v = ws_remove.cell(r, 1).value
        if v:
            text_contents.append(str(v))
    for section_name in REMOVE_SECTIONS:
        assert any(section_name in t for t in text_contents)


def test_job_role_blank_when_not_required():
    responses = excel_bytes([response()])
    placement = excel_bytes([placement_row("R001")])
    _, result = process(responses, placement, open_rules(output={"include_job_role": False}))
    frame = build_rait_frame(result.eligible, result.rules)
    assert frame.iloc[0]["Job Role"] == ""


def test_missing_required_header_fails_validation():
    responses = excel_bytes([{"Timestamp": "2026-01-01", "Full Name": "A"}])
    placement = excel_bytes([placement_row("R001")])
    validation, result = process(responses, placement, open_rules())
    assert result is None
    assert validation.ok is False
    assert any("College Roll No." in err for err in validation.errors)


def test_mobile_number_stored_as_text_and_ignored_errors():
    import zipfile
    from app import clean_download_filename

    responses = excel_bytes([
        response(**{
            "College Roll No.": "R001",
            "Full Name": "Student One",
            "Email address": "student1@college.edu",
            "Mobile No": "9142181780",
        }),
        response(**{
            "College Roll No.": "R002",
            "Full Name": "Student Two",
            "Email address": "student2@college.edu",
            "Mobile No": "09876543210",
            "10th Percentage": 50,  # will go to Remove
        }),
    ])
    placement = excel_bytes([placement_row("R001"), placement_row("R002")])
    rules = open_rules(academic={"mode": "threshold", "tenth": 70, "twelfth": 70, "cgpa": 7.0})
    _, result = process(responses, placement, rules)
    payload = build_workbook(result)


    wb = load_workbook(BytesIO(payload))
    ws_rait = wb["RAIT"]
    # Check column formatting
    assert ws_rait.column_dimensions["E"].number_format == "@"
    # Check cell data type and number format
    assert ws_rait["E2"].data_type == "s"
    assert ws_rait["E2"].number_format == "@"
    assert ws_rait["E2"].value == "9142181780"

    ws_remove = wb["Remove"]
    assert ws_remove.column_dimensions["E"].number_format == "@"

    # Check XML has ignoredErrors for both RAIT and Remove sheets
    z = zipfile.ZipFile(BytesIO(payload))
    sheet1_xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    sheet2_xml = z.read("xl/worksheets/sheet2.xml").decode("utf-8")
    assert "<ignoredErrors>" in sheet1_xml
    assert "numberStoredAsText=\"1\"" in sheet1_xml
    assert "<ignoredErrors>" in sheet2_xml
    assert "numberStoredAsText=\"1\"" in sheet2_xml


def test_clean_download_filename():
    from app import clean_download_filename

    assert clean_download_filename("Deloitte Placement Drive  For Batch 2027 (Responses).xlsx") == "Deloitte Placement Drive For Batch 2027.xlsx"
    assert clean_download_filename("RECRUIT CRM PLACEMENT DRIVE FOR BTECH AND MTECH FOR BATCH 2027 (Responses).xlsx") == "RECRUIT CRM PLACEMENT DRIVE FOR BTECH AND MTECH FOR BATCH 2027.xlsx"
    assert clean_download_filename("Drive (responses).xlsm") == "Drive.xlsm"
    assert clean_download_filename("") == "RAIT_eligible.xlsx"

